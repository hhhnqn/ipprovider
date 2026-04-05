"""Pruebas de extracción unificada (PDF, xlsx, docx)."""

from pathlib import Path

import pytest

from ipprovider.document_extract import extract_text
from ipprovider.ip_find import find_ips_in_text


def _write_pdf_with_text(path: Path, text: str) -> None:
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    c.drawString(72, 800, text)
    c.save()


def test_extract_pdf_via_unified(tmp_path):
    pytest.importorskip("reportlab", reason="pip install -e '.[dev]' para generar PDFs de prueba")
    p = tmp_path / "a.pdf"
    _write_pdf_with_text(p, "Host 203.0.113.5 end")
    t = extract_text(p)
    assert "203.0.113.5" in t
    assert "203.0.113.5" in find_ips_in_text(t)


def test_extract_xlsx(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    p = tmp_path / "w.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "Label"
    ws["B1"] = "Server 198.51.100.10"
    ws.append(["More", "10.0.0.2"])
    wb.save(p)

    t = extract_text(p)
    assert "198.51.100.10" in t
    ips = find_ips_in_text(t)
    assert "198.51.100.10" in ips
    assert "10.0.0.2" in ips


def test_extract_docx(tmp_path):
    pytest.importorskip("docx", reason="pip install -e . (incluye python-docx)")
    from docx import Document

    p = tmp_path / "d.docx"
    doc = Document()
    doc.add_paragraph("Contact at 8.8.8.8 and backup 192.168.0.1")
    doc.save(p)

    t = extract_text(p)
    assert "8.8.8.8" in t
    ips = find_ips_in_text(t)
    assert "8.8.8.8" in ips
    assert "192.168.0.1" in ips


def test_unsupported_xls(tmp_path):
    p = tmp_path / "old.xls"
    p.write_bytes(b"fake")
    with pytest.raises(ValueError, match=r"\.xls"):
        extract_text(p)


def test_unsupported_doc(tmp_path):
    p = tmp_path / "old.doc"
    p.write_bytes(b"fake")
    with pytest.raises(ValueError, match=r"\.doc"):
        extract_text(p)


def test_unsupported_extension(tmp_path):
    p = tmp_path / "f.txt"
    p.write_text("1.1.1.1")
    with pytest.raises(ValueError, match="no soportada"):
        extract_text(p)
