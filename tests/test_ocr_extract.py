"""OCR integrado en extract_text (mocks, sin Tesseract en CI)."""

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from ipprovider.document_extract import extract_text
from ipprovider.ip_find import find_ips_in_text
from ipprovider.ocr_images import iter_ooxml_embedded_images


def _write_pdf_with_text(path: Path, text: str) -> None:
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    c.drawString(72, 800, text)
    c.save()


def test_extract_pdf_with_ocr_merges_mocked_ocr(monkeypatch, tmp_path):
    pytest.importorskip("reportlab")

    p = tmp_path / "a.pdf"
    _write_pdf_with_text(p, "visible 10.0.0.1")

    def fake_ocr(path: Path, *, lang: str, dpi: int = 200, progress: bool = False) -> str:
        assert path == p
        return "OCR only 198.51.100.77"

    monkeypatch.setattr("ipprovider.ocr_images.ocr_pdf_pages", fake_ocr)
    t = extract_text(p, ocr=True)
    assert "10.0.0.1" in t
    assert "198.51.100.77" in t
    assert "198.51.100.77" in find_ips_in_text(t)


def test_extract_xlsx_with_ocr_merges_mocked(monkeypatch, tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    p = tmp_path / "w.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "x"
    wb.save(p)

    def fake_ooxml(path: Path, **kwargs: object) -> str:
        return "203.0.113.10"

    monkeypatch.setattr("ipprovider.ocr_images.ocr_ooxml_embedded_images", fake_ooxml)
    t = extract_text(p, ocr=True)
    assert "203.0.113.10" in t
    assert "203.0.113.10" in find_ips_in_text(t)


def test_extract_docx_with_ocr_merges_mocked(monkeypatch, tmp_path):
    pytest.importorskip("docx")
    from docx import Document

    p = tmp_path / "d.docx"
    doc = Document()
    doc.add_paragraph("hello")
    doc.save(p)

    monkeypatch.setattr(
        "ipprovider.ocr_images.ocr_ooxml_embedded_images",
        lambda path, **kwargs: "2001:db8::1",
    )
    t = extract_text(p, ocr=True)
    assert "2001:db8::1" in find_ips_in_text(t)


def test_iter_ooxml_embedded_images_finds_media(tmp_path):
    zpath = tmp_path / "fake.xlsx"
    png_min = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f\x00"
        b"\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    with ZipFile(zpath, "w", ZIP_DEFLATED) as zf:
        zf.writestr("xl/media/image1.png", png_min)
        zf.writestr("xl/worksheets/sheet1.xml", b"<xml/>")
    blobs = list(iter_ooxml_embedded_images(zpath, "xl/media/"))
    assert len(blobs) == 1
    assert blobs[0] == png_min
