"""Extracción de texto unificada: PDF, Excel (.xlsx) y Word (.docx)."""

from __future__ import annotations

import sys
from pathlib import Path

from ipprovider.pdf_extract import extract_text_from_pdf

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".xlsx", ".docx"})

# Cada cuántas filas mostrar avance en stderr (libros muy grandes).
_XLSX_ROW_PROGRESS_STEP = 5_000

_OCR_SEPARATOR = "\n\n--- Texto OCR (imágenes incrustadas) ---\n\n"


def extract_text(
    path: Path,
    *,
    progress: bool = False,
    ocr: bool = False,
    ocr_lang: str | None = None,
) -> str:
    """
    Extrae texto plano según la extensión del archivo.

    Soporta: .pdf, .xlsx, .docx

    Si ``ocr`` es True, concatena el texto reconocido por OCR en imágenes
    incrustadas (requiere ``pip install -e ".[ocr]"`` y Tesseract; PDF
    además Poppler).
    """
    from ipprovider.ocr_images import DEFAULT_OCR_LANG

    lang = ocr_lang if ocr_lang is not None else DEFAULT_OCR_LANG
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        base = extract_text_from_pdf(path)
        if ocr:
            from ipprovider.ocr_images import ocr_pdf_pages

            if progress:
                print("OCR en páginas PDF (puede tardar)...", file=sys.stderr, flush=True)
            extra = ocr_pdf_pages(path, lang=lang, progress=progress)
            return _merge_ocr(base, extra)
        return base
    if suffix == ".xlsx":
        base = _extract_xlsx(path, progress=progress)
        if ocr:
            from ipprovider.ocr_images import ocr_ooxml_embedded_images

            if progress:
                print("OCR en imágenes incrustadas del Excel...", file=sys.stderr, flush=True)
            extra = ocr_ooxml_embedded_images(
                path,
                media_prefix="xl/media/",
                label="Excel",
                lang=lang,
                progress=progress,
            )
            return _merge_ocr(base, extra)
        return base
    if suffix == ".docx":
        base = _extract_docx(path)
        if ocr:
            from ipprovider.ocr_images import ocr_ooxml_embedded_images

            if progress:
                print("OCR en imágenes incrustadas del Word...", file=sys.stderr, flush=True)
            extra = ocr_ooxml_embedded_images(
                path,
                media_prefix="word/media/",
                label="Word",
                lang=lang,
                progress=progress,
            )
            return _merge_ocr(base, extra)
        return base
    if suffix == ".xls":
        msg = (
            "El formato .xls (Excel 97-2003) no está soportado. "
            "Convierta el archivo a .xlsx o exporte como CSV."
        )
        raise ValueError(msg)
    if suffix == ".doc":
        msg = (
            "El formato .doc (Word binario) no está soportado. "
            "Use .docx o convierta el documento."
        )
        raise ValueError(msg)
    raise ValueError(
        f"Extensión no soportada: {suffix or '(sin extensión)'}. "
        f"Use: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )


def _merge_ocr(base: str, ocr_text: str) -> str:
    ocr_text = (ocr_text or "").strip()
    if not ocr_text:
        return base
    base = base or ""
    if base.strip():
        return base.rstrip() + _OCR_SEPARATOR + ocr_text
    return ocr_text


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _extract_xlsx(path: Path, *, progress: bool = False) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ValueError(
            "Falta el paquete openpyxl (necesario para archivos .xlsx). "
            "Instálelo con: pip install openpyxl   o   pip install -e ."
        ) from e

    if progress:
        print(
            f"Leyendo Excel: {path.name!r} (modo solo lectura)...",
            file=sys.stderr,
            flush=True,
        )

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            if progress:
                print(
                    f"  Hoja «{ws.title}» ...",
                    file=sys.stderr,
                    flush=True,
                )
            row_in_sheet = 0
            for row in ws.iter_rows(values_only=True):
                row_in_sheet += 1
                if progress and row_in_sheet % _XLSX_ROW_PROGRESS_STEP == 0:
                    print(
                        f"    ... {row_in_sheet} filas leídas en «{ws.title}»",
                        file=sys.stderr,
                        flush=True,
                    )
                cells = [_cell_to_str(v) for v in row]
                line = "\t".join(c for c in cells if c)
                if line:
                    parts.append(line)
            parts.append("")
            if progress:
                print(
                    f"  Hoja «{ws.title}» finalizada ({row_in_sheet} filas).",
                    file=sys.stderr,
                    flush=True,
                )
        if progress:
            print("Lectura de Excel terminada.", file=sys.stderr, flush=True)
    finally:
        wb.close()
    return "\n".join(parts)


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise ValueError(
            "Falta el paquete python-docx (necesario para archivos .docx). "
            "Instálelo con: pip install python-docx   o   pip install -e ."
        ) from e

    doc = Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = "\t".join(c for c in cells if c)
            if line:
                parts.append(line)
        parts.append("")
    return "\n".join(parts)
